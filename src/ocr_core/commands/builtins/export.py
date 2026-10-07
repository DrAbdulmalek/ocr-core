from pathlib import Path
from ..core import command, ExecutionContext


@command(
    id="export.markdown", name="تصدير Markdown",
    description="يحوّل النتائج إلى Markdown.",
    category="export",
    params_schema={"type": "object", "properties": {
        "text_key": {"type": "string"}, "output_path": {"type": "string"},
        "title": {"type": "string"}}, "required": ["output_path"]},
)
def export_markdown(ctx: ExecutionContext, text_key: str = "corrected_text",
                     output_path: str = "", title: str = "") -> dict:
    text = ctx.get(text_key, "")
    if not isinstance(text, str):
        text = text.get("text", "") if isinstance(text, dict) else str(text)
    md = f"# {title}\n\n{text}\n" if title else f"{text}\n"
    Path(output_path).write_text(md, encoding="utf-8")
    return {"path": output_path, "length": len(md)}


@command(
    id="export.json", name="تصدير JSON",
    description="يحفظ السياق كـ JSON.",
    category="export",
    params_schema={"type": "object", "properties": {
        "keys": {"type": "array"}, "output_path": {"type": "string"}},
        "required": ["output_path"]},
)
def export_json(ctx: ExecutionContext, keys: list = None, output_path: str = "") -> dict:
    import json
    data = {k: ctx.get(k) for k in (keys or list(ctx.data.keys()))}
    def default(o):
        if hasattr(o, "isoformat"):
            return o.isoformat()
        if hasattr(o, "tobytes"):
            return "<image>"
        return str(o)
    Path(output_path).write_text(
        json.dumps(data, ensure_ascii=False, indent=2, default=default),
        encoding="utf-8")
    return {"path": output_path, "keys": list(data.keys())}


@command(
    id="export.docx", name="تصدير Word (RTL)",
    description="يصدّر النص إلى DOCX بدعم RTL كامل.",
    category="export",
    params_schema={"type": "object", "properties": {
        "text_key": {"type": "string"}, "output_path": {"type": "string"},
        "font_cs": {"type": "string"}}, "required": ["output_path"]},
)
def export_docx(ctx: ExecutionContext, text_key: str = "corrected_text",
                 output_path: str = "",
                 font_cs: str = "Noto Naskh Arabic") -> dict:
    from docx import Document
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    text = ctx.get(text_key, "")
    if not isinstance(text, str):
        text = str(text or "")
    doc = Document()
    for para_text in text.split("\n\n"):
        if not para_text.strip():
            continue
        p = doc.add_paragraph(para_text)
        pPr = p._p.get_or_add_pPr()
        pPr.append(OxmlElement("w:bidi"))
        for run in p.runs:
            rPr = run._r.get_or_add_rPr()
            rPr.append(OxmlElement("w:rtl"))
            fonts = rPr.find(qn("w:rFonts"))
            if fonts is None:
                fonts = OxmlElement("w:rFonts")
                rPr.append(fonts)
            fonts.set(qn("w:cs"), font_cs)
    doc.save(output_path)
    return {"path": output_path, "paragraphs": len(doc.paragraphs)}


@command(
    id="export.hf_dataset", name="تصدير HuggingFace",
    description="يحفظ القصاصات المعتمدة بصيغة HF Datasets.",
    category="export",
    params_schema={"type": "object", "properties": {
        "snippets_key": {"type": "string"}, "output_dir": {"type": "string"},
        "dataset_name": {"type": "string"}}, "required": ["output_dir"]},
)
def export_hf_dataset(ctx: ExecutionContext, snippets_key: str = "snippets",
                       output_dir: str = "",
                       dataset_name: str = "arabic-ocr-snippets") -> dict:
    import json
    from pathlib import Path
    snippets = ctx.get(snippets_key, [])
    out = Path(output_dir)
    (out / "images").mkdir(parents=True, exist_ok=True)
    rows = []
    for i, s in enumerate(snippets):
        if not isinstance(s, dict):
            continue
        img = s.get("image")
        text = s.get("text", "").strip()
        if not text or img is None:
            continue
        img_name = f"snippet_{i:05d}.png"
        try:
            img.save(out / "images" / img_name)
        except Exception:
            continue
        rows.append({"id": i, "image": f"images/{img_name}", "text": text,
                     "language": s.get("language", "ar"),
                     "category": s.get("category", "general"),
                     "level": s.get("level", "word")})
    with open(out / "train.jsonl", "w", encoding="utf-8") as f:
        for r in rows:
            f.write(json.dumps(r, ensure_ascii=False) + "\n")
    (out / "dataset_info.json").write_text(json.dumps({
        "name": dataset_name, "total": len(rows), "language": "ar",
    }, ensure_ascii=False, indent=2), encoding="utf-8")
    return {"path": str(out), "samples": len(rows)}
